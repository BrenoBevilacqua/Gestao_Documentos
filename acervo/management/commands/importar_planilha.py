"""Importa as linhas da aba 'Controle' da planilha Excel para o banco de dados.

Uso:
    python manage.py importar_planilha                       # usa dados/controle_fichas.xlsx
    python manage.py importar_planilha caminho/arquivo.xlsx
"""
from datetime import date, datetime
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from acervo.models import Retirada, Servidor, Solicitante

# Colunas da aba "Controle" (linha de cabeçalho = 3)
COLUNAS = ["Nº", "Nº da Caixa", "Nome do Servidor", "Matrícula", "Solicitante (RH)", "Data de Retirada",
           "Prazo de Devolução", "Data de Devolução", "Situação", "Observações"]


def valor(v):
    """Ignora fórmulas e textos vazios."""
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if not v or v.startswith("="):
            return None
    return v


def para_data(v):
    v = valor(v)
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    texto = str(v).strip()
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%y", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    raise CommandError(f"Data inválida na planilha: {texto!r}")


def para_inteiro(v):
    v = valor(v)
    if v is None:
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def para_texto(v):
    v = valor(v)
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


class Command(BaseCommand):
    help = "Importa as retiradas da aba 'Controle' da planilha de controle de fichas do acervo."

    def add_arguments(self, parser):
        parser.add_argument("arquivo", nargs="?", default=str(Path(settings.BASE_DIR) / "dados" / "controle_fichas.xlsx"))
        parser.add_argument("--aba", default="Controle", help="Nome da aba (padrão: Controle)")

    @transaction.atomic
    def handle(self, arquivo, aba, **options):
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover
            raise CommandError("Instale o openpyxl: pip install openpyxl") from exc

        caminho = Path(arquivo)
        if not caminho.exists():
            raise CommandError(f"Arquivo não encontrado: {caminho}")

        wb = openpyxl.load_workbook(caminho)
        if aba not in wb.sheetnames:
            raise CommandError(f"Aba '{aba}' não encontrada. Abas disponíveis: {', '.join(wb.sheetnames)}")
        ws = wb[aba]

        # Localiza a linha de cabeçalho
        linha_cab = None
        for row in ws.iter_rows(min_row=1, max_row=10):
            textos = [str(c.value).strip() if c.value else "" for c in row]
            if "Nome do Servidor" in textos:
                linha_cab = row[0].row
                idx = {nome: textos.index(nome) for nome in COLUNAS if nome in textos}
                break
        if linha_cab is None:
            raise CommandError("Cabeçalho 'Nome do Servidor' não encontrado na aba.")

        criadas = ignoradas = 0
        for row in ws.iter_rows(min_row=linha_cab + 1, values_only=True):
            get = lambda col: row[idx[col]] if col in idx and idx[col] < len(row) else None  # noqa: E731
            nome = para_texto(get("Nome do Servidor"))
            data_retirada = para_data(get("Data de Retirada"))
            if not nome or not data_retirada:
                continue

            caixa = para_inteiro(get("Nº da Caixa"))
            matricula = para_texto(get("Matrícula"))
            servidor, _ = Servidor.objects.get_or_create(
                nome=nome, matricula=matricula, defaults={"numero_caixa": caixa}
            )
            solicitante, _ = Solicitante.objects.get_or_create(
                nome=para_texto(get("Solicitante (RH)")) or "Não informado", defaults={"setor": "RH"}
            )
            dados = dict(
                servidor=servidor,
                solicitante=solicitante,
                numero_caixa=caixa,
                data_retirada=data_retirada,
            )
            if Retirada.objects.filter(**dados).exists():
                ignoradas += 1
                continue
            Retirada.objects.create(
                **dados,
                prazo_devolucao=para_data(get("Prazo de Devolução")),
                data_devolucao=para_data(get("Data de Devolução")),
                observacoes=para_texto(get("Observações")),
            )
            criadas += 1

        self.stdout.write(self.style.SUCCESS(
            f"Importação concluída: {criadas} retirada(s) criada(s), {ignoradas} já existente(s) ignorada(s)."
        ))
